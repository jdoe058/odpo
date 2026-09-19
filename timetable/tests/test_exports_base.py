"""
Тесты timetable.exports.base.

Проверяют:
* render_docx возвращает валидный .docx с подставленными значениями;
* docx_response формирует корректные заголовки, включая UTF-8 filename;
* тело ответа не искажается.

БД не нужна — SimpleTestCase.
"""
from io import BytesIO
from pathlib import Path
import tempfile

from django.http import HttpResponse
from django.test import SimpleTestCase

from docx import Document

from timetable.exports.base import (
    DOCX_CONTENT_TYPE, docx_response, render_docx,
)
from timetable.tests._factories import make_docx_bytes


class RenderDocxTests(SimpleTestCase):
    """render_docx: docxtpl поверх python-docx."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls._tmp = tempfile.TemporaryDirectory()
        cls.tmp_dir = Path(cls._tmp.name)

    @classmethod
    def tearDownClass(cls):
        cls._tmp.cleanup()
        super().tearDownClass()

    def _write(self, name: str, placeholders: list[str]) -> str:
        path = self.tmp_dir / name
        path.write_bytes(make_docx_bytes(placeholders))
        return str(path)

    def test_returns_bytes(self):
        path = self._write("t.docx", ["name"])
        result = render_docx(path, {"name": "X"})
        self.assertIsInstance(result, bytes)
        self.assertGreater(len(result), 0)

    def test_result_is_valid_docx(self):
        path = self._write("t.docx", ["name"])
        data = render_docx(path, {"name": "ИВАНОВ И.И."})
        doc = Document(BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("ИВАНОВ И.И.", text)
        self.assertNotIn("{{ name }}", text)

    def test_multiple_placeholders(self):
        path = self._write("t.docx", ["a", "b", "c"])
        data = render_docx(path, {"a": "1", "b": "2", "c": "3"})
        doc = Document(BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs)
        for value in ("1", "2", "3"):
            self.assertIn(value, text)

    def test_empty_context(self):
        """Шаблон без плейсхолдеров рендерится пустым контекстом."""
        path = self._write("t.docx", [])
        data = render_docx(path, {})
        self.assertGreater(len(data), 0)

    def test_preserves_original_file(self):
        """render_docx не перезаписывает исходный шаблон."""
        path = self._write("t.docx", ["name"])
        original = Path(path).read_bytes()
        render_docx(path, {"name": "X"})
        self.assertEqual(Path(path).read_bytes(), original)


class DocxResponseTests(SimpleTestCase):
    """docx_response: HTTP-обёртка вокруг байтов .docx."""

    def test_returns_http_response(self):
        response = docx_response(b"x", "f.docx")
        self.assertIsInstance(response, HttpResponse)

    def test_content_type(self):
        response = docx_response(b"x", "f.docx")
        self.assertEqual(response["Content-Type"], DOCX_CONTENT_TYPE)

    def test_content_disposition_ascii(self):
        response = docx_response(b"x", "schedule.docx")
        self.assertEqual(
            response["Content-Disposition"],
            "attachment; filename*=UTF-8''schedule.docx",
        )

    def test_content_disposition_cyrillic_is_encoded(self):
        """
        Кириллица в имени файла должна быть percent-encoded и не
        попадать в заголовок сырыми байтами — иначе HTTP-заголовок
        может быть несовместим с некоторыми клиентами.
        """
        response = docx_response(b"x", "расписание_2026.docx")
        cd = response["Content-Disposition"]
        self.assertTrue(cd.startswith("attachment; filename*=UTF-8''"))
        self.assertNotIn("расписание", cd)
        self.assertIn("%", cd)

    def test_content_disposition_spaces_encoded(self):
        response = docx_response(b"x", "my file.docx")
        cd = response["Content-Disposition"]
        self.assertIn("%20", cd)
        self.assertNotIn("my file.docx", cd)

    def test_body_preserved(self):
        data = b"\x50\x4b\x03\x04payload-not-a-real-docx"
        response = docx_response(data, "f.docx")
        self.assertEqual(response.content, data)

    def test_empty_body_allowed(self):
        """Пустые байты не падают — проверяем, что нет скрытых
        требований к содержимому."""
        response = docx_response(b"", "empty.docx")
        self.assertEqual(response.content, b"")
        self.assertEqual(response["Content-Type"], DOCX_CONTENT_TYPE)