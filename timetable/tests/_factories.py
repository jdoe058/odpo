"""
Хелперы для тестов. Не тесты — общие фабрики.

Импорты моделей — локальные, внутри функций. Причина: `_factories`
может импортироваться на этапе, когда приложения ещё не загружены
(например, из `conftest`-подобных мест). Локальный импорт гарантирует,
что модели доступны только в момент вызова, когда Django уже готов.
"""
from io import BytesIO

from docx import Document


# --- .docx ------------------------------------------------------------

def make_docx_bytes(placeholders: list[str] | None = None) -> bytes:
    """
    Минимальный валидный .docx. Если переданы `placeholders`, каждый
    становится отдельным абзацем вида `{{ name }}` — пригодно для
    проверки рендера через docxtpl.

    Не читает диск и не требует MEDIA_ROOT — возвращает bytes.
    """
    doc = Document()
    for ph in placeholders or []:
        doc.add_paragraph("{{ " + ph + " }}")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --- Модели справочников ---------------------------------------------

def make_position(
    name: str,
    *,
    max_hours_per_day: int = 0,
    max_hours_per_week: int | None = None,
    max_hours_per_year: int | None = None,
    can_sign: bool = False,
    can_approve: bool = False,
    sort_order: int = 100,
):
    from timetable.models import Position
    return Position.objects.create(
        name=name,
        max_hours_per_day=max_hours_per_day,
        max_hours_per_week=max_hours_per_week,
        max_hours_per_year=max_hours_per_year,
        can_sign=can_sign,
        can_approve=can_approve,
        sort_order=sort_order,
    )


def make_employee(short_name: str, position):
    from timetable.models import Employee
    return Employee.objects.create(short_name=short_name, position=position)


def make_base(name: str):
    from timetable.models import Base
    return Base.objects.create(name=name)


def make_funding_type(name: str):
    from timetable.models import FundingType
    return FundingType.objects.create(name=name)


def make_cycle_name(name: str):
    from timetable.models import CycleName
    return CycleName.objects.create(name=name)


def make_cycle(*, name, base, funding, compiled_by, start_date, end_date):
    from timetable.models import Cycle
    return Cycle.objects.create(
        name=name,
        base=base,
        funding_type=funding,
        compiled_by=compiled_by,
        start_date=start_date,
        end_date=end_date,
    )

def make_docx_file(name: str = "t.docx"):
    """ContentFile с минимальным .docx — для FileField."""
    from django.core.files.base import ContentFile
    return ContentFile(make_docx_bytes(), name=name)
