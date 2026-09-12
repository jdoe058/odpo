from datetime import datetime, time
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from timetable.models import (
    Base, Cycle, CycleName, Employee, FundingType, Lesson, LessonType,
)


def parse_time(s: str) -> time:
    """'12.00' / '12:00' -> time(12, 0)."""
    s = str(s).strip().replace(".", ":")
    parts = s.split(":")
    hh = int(parts[0])
    mm = int(parts[1]) if len(parts) > 1 else 0
    return time(hh, mm)


def parse_time_range(value: str):
    """'12.00-13.30' -> (time(12,0), time(13,30))."""
    if not value or "-" not in str(value):
        raise ValueError(f"неверный диапазон: {value!r}")
    a, b = str(value).split("-", 1)
    return parse_time(a), parse_time(b)


def normalize_name(value) -> str:
    """Схлопываем пробелы и в верхний регистр — как в Employee.save()."""
    return " ".join(str(value).split()).upper()


class Command(BaseCommand):
    help = "Импорт расписания из XLSX-файла"

    def add_arguments(self, parser):
        parser.add_argument("file", type=str, help="Путь к XLSX-файлу")
        parser.add_argument(
            "--sheet", type=str, default=None,
            help="Имя листа (по умолчанию — активный)",
        )

    @transaction.atomic
    def handle(self, *args, **options):
        path = Path(options["file"])
        if not path.exists():
            raise CommandError(f"Файл не найден: {path}")

        wb = load_workbook(path, data_only=True)
        ws = wb[options["sheet"]] if options["sheet"] else wb.active

        # --- Шапка ---
        approver_name = normalize_name(ws["A1"].value)
        funding_name = str(ws["A2"].value).strip()
        base_name = str(ws["A3"].value).strip()
        cycle_name_str = str(ws["A4"].value).strip()
        start_date = ws["A5"].value
        end_date = ws["A6"].value

        if isinstance(start_date, datetime):
            start_date = start_date.date()
        if isinstance(end_date, datetime):
            end_date = end_date.date()

        # --- Резолвим справочники ---
        try:
            approver = Employee.objects.get(short_name=approver_name)
        except Employee.DoesNotExist:
            raise CommandError(f"Сотрудник не найден: {approver_name!r}")

        try:
            funding = FundingType.objects.get(name=funding_name)
        except FundingType.DoesNotExist:
            raise CommandError(f"Вид финансирования не найден: {funding_name!r}")

        try:
            base = Base.objects.get(name=base_name)
        except Base.DoesNotExist:
            raise CommandError(f"База не найдена: {base_name!r}")

        try:
            cycle_name = CycleName.objects.get(name=cycle_name_str)
        except CycleName.DoesNotExist:
            raise CommandError(f"Название цикла не найдено: {cycle_name_str!r}")

        cycle, created = Cycle.objects.get_or_create(
            name=cycle_name, start_date=start_date, end_date=end_date,
            defaults={
                "funding_type": funding,
                "base": base,
                "approved_by": approver,
            },
        )
        if created:
            self.stdout.write(self.style.SUCCESS(f"Создан цикл: {cycle}"))
        else:
            self.stdout.write(self.style.WARNING(f"Цикл уже есть: {cycle}"))

        # Перезаписываем занятия этого цикла
        deleted, _ = cycle.lessons.all().delete()
        if deleted:
            self.stdout.write(f"Удалено старых занятий: {deleted}")

        # --- Занятия ---
        created_count = 0
        errors = []

        for row_idx, row in enumerate(
            ws.iter_rows(min_row=7, values_only=True), start=7
        ):
            date_val, time_range, hours, code, topic, teacher = (row + (None,) * 6)[:6]

            if not date_val or not teacher:
                continue

            if isinstance(date_val, datetime):
                date_val = date_val.date()

            try:
                t_start, t_end = parse_time_range(time_range)
            except (ValueError, TypeError):
                errors.append(f"{row_idx}: время {time_range!r}")
                continue

            try:
                hours_int = int(hours)
            except (TypeError, ValueError):
                errors.append(f"{row_idx}: часы {hours!r}")
                continue

            code_str = str(code).strip()
            try:
                lesson_type = LessonType.objects.get(code=code_str)
            except LessonType.DoesNotExist:
                errors.append(f"{row_idx}: LessonType.code={code_str!r} не найден")
                continue

            teacher_name = normalize_name(teacher)
            try:
                teacher_obj = Employee.objects.get(short_name=teacher_name)
            except Employee.DoesNotExist:
                errors.append(f"{row_idx}: сотрудник {teacher_name!r} не найден")
                continue

            Lesson.objects.create(
                cycle=cycle, date=date_val,
                time_start=t_start, time_end=t_end, hours=hours_int,
                lesson_type=lesson_type,
                topic=str(topic or "").strip(),
                employee=teacher_obj,
            )
            created_count += 1

        self.stdout.write(self.style.SUCCESS(
            f"Импортировано занятий: {created_count}"
        ))
        if errors:
            self.stdout.write(self.style.WARNING("Предупреждения:"))
            for e in errors:
                self.stdout.write(f"  - {e}")
