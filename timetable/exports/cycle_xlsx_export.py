"""Выгрузка цикла и его занятий в XLSX."""
import io

from openpyxl import Workbook

from timetable.models import Cycle, Lesson
from timetable.xlsx_cycle_format import (
    SHEET_CYCLE, SHEET_LESSONS,
    set_cycle_sheet_widths, set_lessons_sheet_widths,
    write_cycle_sheet_header, write_lessons_sheet_header,
)

def cycle_to_xlsx_bytes(cycle: Cycle) -> bytes:
    wb = Workbook()

    ws = wb.worksheets[0]
    ws.title = SHEET_CYCLE
    write_cycle_sheet_header(ws,     values = {
        "compiled_by": cycle.compiled_by.short_name,
        "funding": cycle.funding_type.name,
        "base": cycle.base.name,
        "name": cycle.name.name,
        "discipline": (
            cycle.discipline.name_short if cycle.discipline else ""
        ),
        "kind": cycle.kind,
        "stream": cycle.stream,
        "start_date": cycle.start_date.isoformat(),
        "end_date": cycle.end_date.isoformat(),
    })

    set_cycle_sheet_widths(ws)

    ws2 = wb.create_sheet(SHEET_LESSONS)
    write_lessons_sheet_header(ws2)
    ws2.freeze_panes = "A3"

    default_break = Lesson._meta.get_field("break_after_minutes").default

    lessons = (
        Lesson.objects
        .filter(cycle=cycle)
        .select_related("lesson_type", "employee", "base")
        .order_by("date", "time_start", "group")
    )
    prev_date = None
    # available_from всех уже обработанных занятий по дням.
    day_pool: dict[object, list] = {}

    for lesson in lessons:
        break_after = (
            "" if lesson.break_after_minutes == default_break
            else lesson.break_after_minutes
        )

        if lesson.date == prev_date:
            date_cell = ""
        else:
            date_cell = lesson.date.isoformat()
            prev_date = lesson.date

        # «Встык» — если время начала совпадает с максимумом
        # available_from уже прошедших занятий того же дня.
        pool = day_pool.get(lesson.date, [])
        candidates = [v for v in pool if v <= lesson.time_start]
        prev_from = max(candidates) if candidates else None

        if prev_from is not None and lesson.time_start == prev_from:
            time_cell = ""
        else:
            time_cell = lesson.time_start.strftime("%H:%M")

        day_pool.setdefault(lesson.date, []).append(lesson.available_from)

        ws2.append([
            date_cell,
            time_cell,
            lesson.hours,
            lesson.lesson_type.code,
            lesson.category or "",
            lesson.group or "",
            lesson.topic or "",
            lesson.employee.short_name,
            break_after,
            _base_name(lesson),
        ])

    set_lessons_sheet_widths(ws2)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

def _base_name(lesson) -> str:
    base = getattr(lesson, "base", None)
    return getattr(base, "name", "") if base is not None else ""

def build_cycle_import_template() -> bytes:
    wb = Workbook()

    ws = wb.worksheets[0]
    ws.title = SHEET_CYCLE
    write_cycle_sheet_header(ws)
    set_cycle_sheet_widths(ws)

    ws2 = wb.create_sheet(SHEET_LESSONS)
    write_lessons_sheet_header(ws2)
    ws2.append([
        "2026-08-24", "09:00", 2, "0", "Основы охраны труда",
        "МАРКОВ Д.В.", 10, "",
    ])
    set_lessons_sheet_widths(ws2)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()