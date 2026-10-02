"""Выгрузка отчёта «Распределение часов за месяц» в XLSX."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from timetable.services.ped_hours import MONTH_NAMES_RU, PedHoursReport


HEADER = (
    "Пед. часы", "ДПП", "Дата начала", "Дата окончания",
    "ФИО", "База", "Обучение",
)

COLUMN_WIDTHS = (10, 60, 14, 14, 22, 40, 14)


def ped_hours_to_xlsx_bytes(report: PedHoursReport) -> bytes:
    wb = Workbook()
    ws = wb.worksheets[0]
    ws.title = "Пед. часы"

    month_label = f"{MONTH_NAMES_RU[report.month]} {report.year}"
    title = f"Распределение часов за {month_label} г."
    summary = (
        f"Всего часов: {report.total_hours}, "
        f"Всего расписаний: {report.total_cycles}"
    )

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADER))
    cell = ws.cell(row=1, column=1, value=title)
    cell.font = Font(bold=True, size=14)
    cell.alignment = Alignment(horizontal="center")

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(HEADER))
    cell = ws.cell(row=2, column=1, value=summary)
    cell.alignment = Alignment(horizontal="center")

    header_row = 3
    for i, name in enumerate(HEADER, start=1):
        c = ws.cell(row=header_row, column=i, value=name)
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r_idx, row in enumerate(report.rows, start=header_row + 1):
        ws.cell(row=r_idx, column=1, value=row.hours)
        ws.cell(row=r_idx, column=2, value=row.cycle_name)
        ws.cell(row=r_idx, column=3, value=row.start_date.strftime("%d.%m.%Y"))
        ws.cell(row=r_idx, column=4, value=row.end_date.strftime("%d.%m.%Y"))
        ws.cell(row=r_idx, column=5, value=row.compiled_by)
        ws.cell(row=r_idx, column=6, value=row.base)
        ws.cell(row=r_idx, column=7, value=row.funding)

    totals_row = header_row + len(report.rows) + 2
    ws.cell(row=totals_row, column=1, value="Итоги по заведующим").font = Font(bold=True)
    for i, t in enumerate(report.totals, start=totals_row + 1):
        ws.cell(row=i, column=1, value=t.compiled_by)
        ws.cell(row=i, column=2, value=t.hours)

    for i, width in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()