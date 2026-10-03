"""Импорт справочников из XLSX.

Один лист = один справочник. Имя листа ищется по `title` из реестра.
Ряд 1 — русские заголовки (для человека), ряд 2 — латиница (для парсера),
данные с третьего ряда. Колонки матчатся строго по латинице.
"""
from dataclasses import dataclass, field
from io import BytesIO
from typing import Any

from django.db import transaction
from openpyxl import load_workbook

from .registry import Column, specs_in_import_order


HEADER_ROW_RU = 1
HEADER_ROW_LATIN = 2
DATA_START_ROW = 3


@dataclass
class ParseResult:
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    stats: dict[str, int] = field(default_factory=dict)
    rows: dict[str, list[dict]] = field(default_factory=dict)

    @property
    def has_errors(self) -> bool:
        return bool(self.errors)


def parse_and_validate(content: bytes) -> ParseResult:
    result = ParseResult()

    try:
        wb = load_workbook(BytesIO(content), data_only=True)
    except Exception as exc:
        result.errors.append(f"Не удалось открыть файл: {exc}")
        return result

    known_titles = {spec.title for spec in specs_in_import_order()}
    for sheet_name in wb.sheetnames:
        if sheet_name not in known_titles:
            result.warnings.append(
                f"лист «{sheet_name}» не распознан — пропущен"
            )

    for spec in specs_in_import_order():
        if spec.title not in wb.sheetnames:
            result.warnings.append(
                f"лист «{spec.title}» отсутствует — пропущен"
            )
            continue
        _parse_sheet(wb[spec.title], spec, result)

    return result


def apply(content: bytes) -> ParseResult:
    result = parse_and_validate(content)
    if result.errors:
        return result

    with transaction.atomic():
        for spec in specs_in_import_order():
            rows = result.rows.get(spec.slug, [])
            for row in rows:
                obj = spec.model.objects.get(pk=row["_pk"])
                for key, value in row.items():
                    if key == "_pk":
                        continue
                    setattr(obj, key, value)
                obj.save()
            if rows:
                result.stats[spec.slug] = len(rows)

    return result


# --- Внутреннее ---------------------------------------------------------

def _parse_sheet(ws, spec, result: ParseResult) -> None:
    by_name = {col.name: col for col in spec.columns}
    columns: list[Column | None] = []
    matched: set[str] = set()

    for idx in range(1, ws.max_column + 1):
        raw = ws.cell(row=HEADER_ROW_LATIN, column=idx).value
        name = str(raw).strip() if raw is not None else ""
        if not name:
            columns.append(None)
            continue
        col = by_name.get(name)
        if col is None:
            result.warnings.append(
                f"[{spec.title}] неизвестная колонка «{name}» — пропущена"
            )
            columns.append(None)
            continue
        columns.append(col)
        matched.add(name)

    missing = [c.name for c in spec.columns if c.name not in matched]
    if missing:
        result.errors.append(
            f"[{spec.title}] в шапке нет колонок: {', '.join(missing)}"
        )
        return

    for row_idx in range(DATA_START_ROW, ws.max_row + 1):
        values = [
            ws.cell(row=row_idx, column=idx).value
            for idx in range(1, ws.max_column + 1)
        ]
        if all(_is_empty(v) for v in values):
            continue
        _parse_row(row_idx, values, columns, spec, result)


def _parse_row(row_idx, values, columns, spec, result: ParseResult) -> None:
    row: dict[str, Any] = {}

    for col, raw in zip(columns, values):
        if col is None:
            continue
        try:
            row[col.name] = _convert(col, raw)
        except ValueError as exc:
            result.warnings.append(
                f"[{spec.title}], строка {row_idx}: поле «{col.name}»: "
                f"{exc} — строка пропущена"
            )
            return

    if not _resolve_foreign_keys(row, row_idx, spec, columns, result):
        return
    if not _resolve_upsert_key(row, row_idx, spec, result):
        return

    result.rows.setdefault(spec.slug, []).append(row)


def _resolve_foreign_keys(row, row_idx, spec, columns, result: ParseResult) -> bool:
    for col in columns:
        if col is None or col.kind != "fk_name":
            continue
        value = row.get(col.name)
        if not value:
            field = spec.model._meta.get_field(col.name)
            if field.null:
                del row[col.name]
                row[col.name + "_id"] = None
                continue
            result.errors.append(
                f"[{spec.title}], строка {row_idx}: поле «{col.name}» пусто"
            )
            return False
        related_model = spec.model._meta.get_field(col.name).related_model
        related = related_model.objects.filter(**{col.fk_attr: value}).first()
        if related is None:
            result.warnings.append(
                f"[{spec.title}], строка {row_idx}: {col.name}=«{value}» "
                "не найдено — строка пропущена"
            )
            return False
        del row[col.name]
        row[col.name + "_id"] = related.pk
    return True


def _resolve_upsert_key(row, row_idx, spec, result: ParseResult) -> bool:
    keys = (
        spec.upsert_key
        if isinstance(spec.upsert_key, tuple)
        else (spec.upsert_key,)
    )

    lookup: dict = {}
    for key in keys:
        # FK-поле уже разрешилось в `<key>_id` в предыдущем шаге.
        if f"{key}_id" in row:
            lookup[f"{key}_id"] = row[f"{key}_id"]
        else:
            lookup[key] = row.get(key)

    # Нормализация — только для одиночного текстового ключа.
    if isinstance(spec.upsert_key, str) and spec.key_normalizer is not None:
        value = lookup.get(spec.upsert_key)
        if value is not None:
            value = spec.key_normalizer(value)
            lookup[spec.upsert_key] = value
            row[spec.upsert_key] = value

    missing = [k for k, v in lookup.items() if v in (None, "")]
    if missing:
        result.errors.append(
            f"[{spec.title}], строка {row_idx}: пустой ключ "
            f"({', '.join(missing)})"
        )
        return False

    obj = spec.model.objects.filter(**lookup).first()
    if obj is None:
        result.errors.append(
            f"[{spec.title}], строка {row_idx}: запись по ключу "
            f"{lookup} не найдена"
        )
        return False

    row["_pk"] = obj.pk
    return True

# --- Типы ---------------------------------------------------------------

def _is_empty(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _convert(col: Column, raw):
    if col.kind == "str":
        return _as_str(raw)
    if col.kind == "int":
        return _as_int(raw)
    if col.kind == "int_nullable":
        return None if _is_empty(raw) else _as_int(raw)
    if col.kind == "bool":
        return _as_bool(raw)
    if col.kind == "fk_name":
        return _as_str(raw)
    raise ValueError(f"неизвестный тип колонки {col.kind!r}")


def _as_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, str):
        return v.strip()
    return str(v).strip()


def _as_int(v) -> int:
    if _is_empty(v):
        raise ValueError("пустое значение")
    if isinstance(v, bool):
        raise ValueError("ожидалось целое число")
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        if v != int(v):
            raise ValueError(f"не целое число: {v}")
        return int(v)
    if isinstance(v, str):
        try:
            return int(v.strip())
        except ValueError:
            raise ValueError(f"не число: {v!r}")
    raise ValueError(f"не число: {v!r}")


def _as_bool(v) -> bool:
    if v is None:
        return False
    if isinstance(v, bool):
        return v
    if isinstance(v, str):
        low = v.strip().lower()
        if low in ("true", "1", "да", "yes"):
            return True
        if low in ("false", "0", "нет", "no", ""):
            return False
        raise ValueError(f"не удалось разобрать boolean: {v!r}")
    if isinstance(v, (int, float)):
        return bool(v)
    raise ValueError(f"не удалось разобрать boolean: {v!r}")