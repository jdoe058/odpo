"""Выгрузка всех справочников в один XLSX.

Один лист = один справочник. Имя листа — человекочитаемое (`title`).
Строка 1 — русские заголовки.
Строка 2 — латинские имена полей (как в модели).
Данные — с третьей строки.
"""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from .registry import Column, specs_in_import_order


HEADER_ROW_RU = 1
HEADER_ROW_LATIN = 2
DATA_START_ROW = 3


def references_to_xlsx_bytes() -> bytes:
    wb = Workbook()
    wb.remove(wb.active)   # дефолтный "Sheet" не нужен

    for spec in specs_in_import_order():
        ws = wb.create_sheet(title=spec.title)
        _write_headers(ws, spec.columns)
        _write_rows(ws, spec)
        _autosize_columns(ws, spec.columns)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _write_headers(ws, columns: tuple[Column, ...]) -> None:
    for idx, col in enumerate(columns, start=1):
        ru = ws.cell(row=HEADER_ROW_RU, column=idx, value=col.header or col.name)
        ru.font = Font(bold=True)
        ru.alignment = Alignment(horizontal="center", wrap_text=True)

        latin = ws.cell(row=HEADER_ROW_LATIN, column=idx, value=col.name)
        latin.font = Font(italic=True, color="777777")
        latin.alignment = Alignment(horizontal="center")


def _write_rows(ws, spec) -> None:
    objects = spec.model.objects.order_by("pk")
    for row_idx, obj in enumerate(objects, start=DATA_START_ROW):
        for col_idx, col in enumerate(spec.columns, start=1):
            ws.cell(row=row_idx, column=col_idx, value=_serialize(obj, col))


def _autosize_columns(ws, columns: tuple[Column, ...]) -> None:
    for idx, col in enumerate(columns, start=1):
        letter = ws.cell(row=1, column=idx).column_letter
        title = col.header or col.name
        ws.column_dimensions[letter].width = max(14, len(title) + 2)


def _serialize(obj, col: Column):
    if col.kind == "fk_name":
        related = getattr(obj, col.name, None)
        if related is None:
            return ""
        return str(getattr(related, col.fk_attr, ""))

    value = getattr(obj, col.name)

    if value is None:
        return ""

    if col.kind == "bool":
        return bool(value)   # openpyxl запишет как boolean-ячейку

    return value