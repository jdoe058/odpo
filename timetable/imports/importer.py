from dataclasses import dataclass
from datetime import date, time
from io import BytesIO

from django.db import transaction
from openpyxl import load_workbook

from timetable.models import (
    Base, Cycle, CycleName, Employee, FundingType, Lesson, LessonType,
    normalize_name, normalize_short_name,
)
from timetable.imports.errors import ScheduleImportError, group_errors
from timetable.imports.parsers import as_date, parse_range


@dataclass
class ParsedLesson:
    row_idx: int
    date: date
    time_start: time
    time_end: time
    hours: int
    lesson_type: LessonType
    topic: str
    employee: Employee


@dataclass
class ImportResult:
    cycle: Cycle
    lessons_created: int


def import_schedule(file_obj) -> ImportResult:
    try:
        wb = load_workbook(BytesIO(file_obj.read()), data_only=True)
    except Exception as e:
        raise ScheduleImportError([f"Не удалось открыть файл: {e}"])

    ws = wb.active
    errors: list[str] = []

    # Фаза 1: шапка
    header, header_errors = parse_header(ws)
    errors.extend(header_errors)

    # Фаза 2: занятия
    parsed, lesson_errors = parse_lessons(ws)
    errors.extend(lesson_errors)

    # Фаза 3: если есть ошибки — ни одной записи
    if errors:
        raise ScheduleImportError(group_errors(errors))

    # Фаза 4: пишем всё одной транзакцией
    return write_cycle(header, parsed)

# ФАЗЫ

@dataclass
class Header:
    cycle_name: CycleName
    funding: FundingType
    base: Base
    compiled_by: Employee
    start_date: date
    end_date: date


def parse_header(ws) -> tuple[Header | None, list[str]]:
    errors: list[str] = []

    compiled_name = normalize_short_name(ws["A1"].value)
    funding_name = normalize_name(ws["A2"].value)
    base_name = normalize_name(ws["A3"].value)
    cycle_name_str = normalize_name(ws["A4"].value)

    try:
        start_date = as_date(ws["A5"].value, row_idx=5)
    except ValueError as e:
        errors.append(str(e))
        start_date = None
    try:
        end_date = as_date(ws["A6"].value, row_idx=6)
    except ValueError as e:
        errors.append(str(e))
        end_date = None

    compiled_by = Employee.objects.filter(short_name=compiled_name).first()
    if compiled_by is None:
        errors.append(f"Сотрудник не найден: {compiled_name!r}")

    funding = FundingType.objects.filter(name=funding_name).first()
    if funding is None:
        errors.append(f"Вид финансирования не найден: {funding_name!r}")

    base = Base.objects.filter(name=base_name).first()
    if base is None:
        errors.append(f"База не найдена: {base_name!r}")

    cycle_name = CycleName.objects.filter(name=cycle_name_str).first()
    if cycle_name is None:
        errors.append(f"Название цикла не найдено: {cycle_name_str!r}")

    if errors:
        return None, errors

    return Header(
        cycle_name=cycle_name, funding=funding, base=base,
        compiled_by=compiled_by, start_date=start_date, end_date=end_date,
    ), []


def parse_lessons(ws) -> tuple[list[ParsedLesson], list[str]]:
    parsed: list[ParsedLesson] = []
    errors: list[str] = []

    cache_types = {lt.code: lt for lt in LessonType.objects.all()}
    cache_employees = {e.short_name: e for e in Employee.objects.all()}

    for row_idx, row in enumerate(
        ws.iter_rows(min_row=7, values_only=True), start=7
    ):
        row = (row + (None,) * 6)[:6]
        date_val, time_range, hours, code, topic, teacher = row

        if not date_val and not teacher:
            continue

        try:
            lesson_date = as_date(date_val, row_idx)
        except ValueError as e:
            errors.append(str(e))
            continue

        try:
            t_start, t_end = parse_range(time_range, row_idx)
        except (ValueError, TypeError) as e:
            errors.append(str(e))
            continue

        try:
            hours_int = int(hours)
        except (TypeError, ValueError):
            errors.append(f"Строка {row_idx}: не распознаны часы {hours!r}")
            continue

        code_str = str(code).strip()
        lesson_type = cache_types.get(code_str)
        if lesson_type is None:
            errors.append(f"Строка {row_idx}: тип занятия с кодом {code_str!r} не найден")
            continue

        teacher_name = normalize_short_name(teacher)
        teacher_obj = cache_employees.get(teacher_name)
        if teacher_obj is None:
            errors.append(f"Строка {row_idx}: сотрудник {teacher_name!r} не найден")
            continue

        parsed.append(ParsedLesson(
            row_idx=row_idx, date=lesson_date,
            time_start=t_start, time_end=t_end,
            hours=hours_int, lesson_type=lesson_type,
            topic=str(topic or "").strip(), employee=teacher_obj,
        ))

    return parsed, errors


def write_cycle(header: Header, parsed: list[ParsedLesson]) -> ImportResult:
    with transaction.atomic():
        cycle, _ = Cycle.objects.get_or_create(
            name=header.cycle_name,
            start_date=header.start_date,
            end_date=header.end_date,
            defaults={
                "funding_type": header.funding,
                "base": header.base,
                "compiled_by": header.compiled_by,
            },
        )

        if cycle.in_archive:
            raise ScheduleImportError([
                f"Цикл «{cycle}» помечен как архивный — импорт запрещён. "
                "Снимите флаг «В архиве» в админке, если нужно переимпортировать."
            ])

        cycle.lessons.all().delete()

        Lesson.objects.bulk_create([
            Lesson(
                cycle=cycle, date=p.date,
                time_start=p.time_start, time_end=p.time_end,
                hours=p.hours, lesson_type=p.lesson_type,
                topic=p.topic, employee=p.employee,
            )
            for p in parsed
        ])

    return ImportResult(cycle=cycle, lessons_created=len(parsed))

