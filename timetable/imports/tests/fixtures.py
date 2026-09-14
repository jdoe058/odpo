from io import BytesIO
from openpyxl import Workbook


def make_xlsx(header: dict, rows: list[tuple]) -> BytesIO:
    """
    header: {'A1': ..., 'A2': ..., ..., 'A6': ...}
    rows: список кортежей (date, time_range, hours, code, topic, teacher)
    """
    wb = Workbook()
    ws = wb.active

    for cell, value in header.items():
        ws[cell] = value

    for r, row in enumerate(rows, start=7):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

