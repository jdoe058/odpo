"""
Разбор XLSX до обращения к БД: какие сущности упомянуты в файле.
"""
from dataclasses import dataclass, field
from io import BytesIO

from openpyxl import load_workbook

from timetable.models import normalize_name, normalize_short_name


@dataclass
class DiscoveredReferences:
    funding_type: str = ""
    base: str = ""
    cycle_name: str = ""
    lesson_types: set[str] = field(default_factory=set)
    employees: set[str] = field(default_factory=set)


def discover_references(data: bytes) -> DiscoveredReferences:
    """
    Открывает XLSX и возвращает упомянутые в нём сущности.

    К БД не обращается, ошибок не группирует — только собирает имена.
    Составитель цикла (A1) попадает в общий набор employees: для БД
    это такой же Employee, как и преподаватели в строках.
    """
    wb = load_workbook(BytesIO(data), data_only=True)
    ws = wb.active

    refs = DiscoveredReferences(
        funding_type=normalize_name(ws["A2"].value),
        base=normalize_name(ws["A3"].value),
        cycle_name=normalize_name(ws["A4"].value),
    )

    compiled_by = normalize_short_name(ws["A1"].value)
    if compiled_by:
        refs.employees.add(compiled_by)

    for row in ws.iter_rows(min_row=7, values_only=True):
        row = (row + (None,) * 6)[:6]
        _date, _time_range, _hours, code, _topic, teacher = row

        if not _date and not teacher:
            continue

        code_str = str(code or "").strip()
        if code_str:
            refs.lesson_types.add(code_str)

        teacher_name = normalize_short_name(teacher)
        if teacher_name:
            refs.employees.add(teacher_name)

    return refs